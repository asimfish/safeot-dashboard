from pathlib import Path
import json,sys,copy,shutil,hashlib,zipfile
import fitz
from pptx import Presentation
from pptx.util import Pt
from PIL import Image
from lxml import etree
import numpy as np
import math_tools

HERE=Path(__file__).resolve().parent; R=HERE.parent
sys.path.append(str(R.parent/'image_pipeline_20260925/runtime-venv/lib/python3.12/site-packages'))
from super_img2ppt.render import LibreOfficeRenderer,rasterize_pdf

def export(n,build,iteration):
 d=HERE/f'conversion/figure{n}/redrawn';src=d/build
 v=json.loads((src/'validation.json').read_text());assert v['status'] in ('pass','review'),v['status']
 out=HERE/f'native/{iteration}/figure{n}';out.mkdir(parents=True,exist_ok=False)
 # Runtime skeleton is the baseline used for the semantic reconstruction;
 # exact formula content is added as stable LaTeX outlines, never raster text.
 files=list(src.glob('*.pptx'));assert len(files)==1
 prs=Presentation(files[0]);sl=prs.slides[0]
 rows=json.loads((d/'equations.json').read_text());math_tools.OUT=out
 md=math_tools.typeset(rows,n);doc=fitz.open(md/'equations.pdf');audit=[]
 svgs=sorted(md.glob('eq-*.svg'),key=lambda p:int(p.stem.split('-')[-1]))
 assert len(svgs)==len(rows)==len(doc)
 for row,page,svg in zip(rows,doc,svgs):
  name,tex,b,size,c=row
  a=math_tools.import_formula(sl,svg,name,b,c,page.rect)
  a.update(latex=tex,minimum_math_pt=min(s['size'] for bl in page.get_text('dict')['blocks'] for l in bl.get('lines',[]) for s in l['spans'] if s['text'].strip()))
  audit.append(a)
 assert all(a['fits'] for a in audit),[a for a in audit if not a['fits']]
 # One editable group per large module, with individual text/shape and formula
 # children. External connectors stay independent and do not auto-reroute.
 groups=['penalty','safeot','constraints','ours_band'] if n==1 else ['rollout','inputs','budgets','projection','update']
 group_report=[]
 for name in groups:
  anchor=next(s for s in sl.shapes if s.name==name)
  x,y,w,h=anchor.left,anchor.top,anchor.width,anchor.height
  members=[s for s in sl.shapes if s.name==name or (s.left>=x and s.top>=y and s.left+s.width<=x+w and s.top+s.height<=y+h)]
  members=[anchor]+[s for s in members if s.name!=name]
  index=list(sl.shapes._spTree).index(anchor._element)
  g=sl.shapes.add_group_shape();g.name='module_'+name
  for s in members:g._element.append(s._element)
  g._element.recalculate_extents();sl.shapes._spTree.remove(g._element);sl.shapes._spTree.insert(index,g._element)
  group_report.append({'id':g.name,'members':[s.name for s in members]})
 if n==2:
  # These boundary-crossing connectors must paint over the lower band's fill.
  for name in ('price_to_sum_1','gae_drop_0','observed_feedback_2','observed_feedback_3'):
   sh=next((s for s in sl.shapes if s.name==name),None)
   if sh is not None:
    sl.shapes._spTree.remove(sh._element);sl.shapes._spTree.append(sh._element)
 props=prs.core_properties
 props.author='';props.last_modified_by='';props.comments='';props.title=f'SafeOT Figure {n}';props.subject='';props.keywords=''
 # Submission notes contain no process URLs, host names or source venue names.
 if sl.has_notes_slide:sl.notes_slide.notes_text_frame.text=''
 stem=f'Figure{n}_StyleCurrent_R1';ppt=out/(stem+'.pptx');prs.save(ppt)
 pdf=LibreOfficeRenderer('/usr/bin/soffice',timeout=60).render(ppt,out/'render')
 doc=fitz.open(pdf);doc.set_metadata({});doc.del_xml_metadata();doc.save(out/(stem+'.pdf'),garbage=4,deflate=True)
 page=doc[0];page.get_pixmap(dpi=300).save(out/(stem+'_300dpi.png'))
 (out/(stem+'.svg')).write_text(page.get_svg_image(text_as_path=False))
 Image.open(out/(stem+'_300dpi.png')).convert('L').save(out/(stem+'_grayscale.png'),dpi=(300,300))
 im=np.asarray(Image.open(out/(stem+'_300dpi.png')).convert('RGB'),dtype=float)/255
 linear=np.where(im<=.04045,im/12.92,((im+.055)/1.055)**2.4)
 for name,m in {'deuteranopia':[[.367322,.860646,-.227968],[.280085,.672501,.047413],[-.011820,.042940,.968881]],'protanopia':[[.152286,1.052583,-.204868],[.114503,.786281,.099216],[-.003882,-.048116,1.051998]]}.items():
  c=np.clip(linear@np.array(m).T,0,1);rgb=np.where(c<=.0031308,12.92*c,1.055*c**(1/2.4)-.055)
  Image.fromarray(np.round(rgb*255).astype('uint8')).save(out/(stem+'_'+name+'.png'),dpi=(300,300))
 spans=[s for b in page.get_text('dict')['blocks'] for l in b.get('lines',[]) for s in l['spans'] if s['text'].strip()]
 with zipfile.ZipFile(ppt) as z:
  xml=etree.fromstring(z.read('ppt/slides/slide1.xml'));ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
  pics=xml.findall('.//p:pic',ns);ids=[x.get('id') for x in xml.findall('.//p:cNvPr',ns)]
  assert len(ids)==len(set(ids));assert not pics
 mintext=min(s['size'] for s in spans);minmath=min(a['minimum_math_pt'] for a in audit)
 assert mintext>=8 and minmath>=8
 assert not page.get_images();assert all(fitz.Rect(s['bbox']) in page.rect for s in spans)
 report={'figure':n,'source_candidate':'New style/current-content S5 F02; not the earlier selected F02','status':'review candidate, no cold-reader acceptance or manuscript promotion','width_in':page.rect.width/72,'height_in':page.rect.height/72,'minimum_pdf_text_pt':mintext,'minimum_math_source_pt':minmath,'math_measurement':'All accents and sub/superscripts in pdflatex/newtx source PDF at scale 1. Final math is editable vector contours, not semantic OfficeMath.','pptx_rasters':len(pics),'pdf_rasters':len(page.get_images()),'all_formula_boxes_fit':all(a['fits'] for a in audit),'renderer':'Actual PPTX exported by LibreOffice through super_img2ppt','powerpoint_verified':False,'source_scene_skeleton_validation':v['status'],'visual_review':'required, final render inspected separately','groups':group_report,'native_shapes':len(xml.findall('.//p:sp',ns))}
 (out/'font_editability_report.json').write_text(json.dumps(report,indent=2));(out/'formula_audit.json').write_text(json.dumps(audit,indent=2))
 # A reversible editability probe on a disposable deck, original is untouched.
 probe=Presentation(ppt);g=next(s for s in probe.slides[0].shapes if s.name=='module_'+('safeot' if n==1 else 'projection'))
 before={s.name:[s.left,s.top,s.width,s.height] for s in probe.slides[0].shapes};g.left+=Pt(2)
 p=out/'editability_probe.pptx';probe.save(p);after=Presentation(p)
 checks={s.name:[s.left,s.top,s.width,s.height] for s in after.slides[0].shapes}
 assert all(before[k]==checks[k] for k in before if k!=g.name)
 (out/'editability_probe.json').write_text(json.dumps({'module_moved_pt':2,'other_top_level_objects_unchanged':True,'source_unchanged':True},indent=2))
 print(json.dumps({k:v for k,v in report.items() if k not in ('groups',)},indent=2))

if __name__=='__main__':export(int(sys.argv[1]),sys.argv[2],sys.argv[3])
