from pathlib import Path

import copy,json,os,re,subprocess,sys,hashlib

from lxml import etree

from pptx import Presentation

from pptx.util import Pt

from pptx.dml.color import RGBColor

from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR

from pptx.enum.text import PP_ALIGN,MSO_ANCHOR

from pptx.oxml.xmlchemy import OxmlElement

from fontTools.svgLib.path import parse_path

from fontTools.pens.recordingPen import RecordingPen

from fontTools.pens.transformPen import TransformPen

from fontTools.pens.boundsPen import BoundsPen

from fontTools.misc.transform import Transform

import fitz

HERE=Path(__file__).resolve().parent

OUT=HERE/"math"; OUT.mkdir(exist_ok=True)

ANS="http://schemas.openxmlformats.org/drawingml/2006/main"

ENV={**os.environ,"PATH":"/usr/bin:/bin:"+os.environ.get("PATH","")}

def clean(s):
 for r in s._element.findall('.//{'+ANS+'}effectRef'):r.set('idx','0')

def typeset(rows,n):
 d=OUT/f'figure{n}';d.mkdir(exist_ok=True)
 header=r'''\documentclass{article}
\usepackage{newtxtext,newtxmath}
\pagestyle{empty}
'''
 for s in sorted(set(r[3] for r in rows)):header+=rf'\DeclareMathSizes{{{s}}}{{{s}}}{{8.2}}{{8.2}}'+'\n'
 header+='\\begin{document}\n'
 for name,tex,b,size,color in rows:
  header+=rf'''\begingroup
\fontsize{{{size}}}{{14}}\selectfont
\setbox0\hbox{{$ {tex} $}}
\pdfpagewidth=\dimexpr\wd0+2pt\relax
\pdfpageheight=\dimexpr\ht0+\dp0+2pt\relax
\pdfhorigin=0pt \pdfvorigin=0pt
\shipout\vbox{{\offinterlineskip\kern1pt\hbox{{\kern1pt\box0\kern1pt}}\kern1pt}}
\endgroup
'''
 header+='\\end{document}\n';(d/'equations.tex').write_text(header)
 p=subprocess.run(['/usr/bin/pdflatex','-no-shell-escape','-interaction=batchmode','-halt-on-error','equations.tex'],cwd=d,env=ENV,capture_output=True,text=True)
 (d/'tex_build.log').write_text(p.stdout+p.stderr);assert p.returncode==0,(d/'equations.log').read_text()[-3000:]
 p=subprocess.run(['/usr/bin/dvisvgm','--pdf','--no-fonts','--exact','--page=1-','--output=eq-%p.svg','equations.pdf'],cwd=d,capture_output=True,text=True)
 (d/'svg_build.log').write_text(p.stdout+p.stderr);assert p.returncode==0,p.stderr
 return d

def import_formula(sl,svg,name,box,color,page_rect):
 root=etree.parse(str(svg)).getroot()
 # dvisvgm 2.13 accumulates the max media box across pages. Each page's
 # original PDF rectangle gives the correct origin; glyph coordinates are unchanged.
 vb=[0,-page_rect.height,page_rect.width,page_rect.height];vw,vh=vb[2:]
 x,y,w,h=box;ox=x+(w-vw)/2-vb[0];oy=y+(h-vh)/2-vb[1]
 if vw>w+.2 or vh>h+.2:print('FORMULA_OVERFLOW',name,vw,vh,box)
 g=sl.shapes.add_group_shape();g.name=name
 paths=[]
 def walk(node,tr=Transform()):
  val=node.get('transform','')
  if val:
   nums=list(map(float,re.findall(r'[-+]?(?:\d*\.)?\d+(?:[eE][-+]?\d+)?',val)));assert val.startswith('matrix('),val
   tr=tr.transform(Transform(*nums))
  if node.tag.endswith('}path'):
   rec=RecordingPen();parse_path(node.get('d'),TransformPen(rec,tr.translate(0,0)))
   pts=BoundsPen(None);rec.replay(pts)
   if pts.bounds:paths.append((rec,pts.bounds))
  for ch in node:walk(ch,tr)
 walk(root)
 for idx,(rec,(xmin,ymin,xmax,ymax)) in enumerate(paths):
  ew=max(.02,xmax-xmin);eh=max(.02,ymax-ymin)
  sh=g.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(ox+xmin),Pt(oy+ymin),Pt(ew),Pt(eh));sh.name=f'{name}_outline_{idx}'
  sp=sh._element.spPr
  for z in list(sp):
   if z.tag.endswith('prstGeom'):sp.remove(z)
  geom=OxmlElement('a:custGeom')
  for tag in ('avLst','gdLst','ahLst','cxnLst'):geom.append(OxmlElement('a:'+tag))
  rect=OxmlElement('a:rect')
  for k,v in dict(l='0',t='0',r='r',b='b').items():rect.set(k,v)
  geom.append(rect);pl=OxmlElement('a:pathLst');path=OxmlElement('a:path');path.set('w',str(round(ew*12700)));path.set('h',str(round(eh*12700)))
  for op,points in rec.value:
   if op in ('closePath','endPath'):
    if op=='closePath':path.append(OxmlElement('a:close'))
    continue
   tag={'moveTo':'moveTo','lineTo':'lnTo','curveTo':'cubicBezTo','qCurveTo':'quadBezTo'}[op]
   node=OxmlElement('a:'+tag)
   for xx,yy in points:
    pt=OxmlElement('a:pt');pt.set('x',str(round((xx-xmin)*12700)));pt.set('y',str(round((yy-ymin)*12700)));node.append(pt)
   path.append(node)
  pl.append(path);geom.append(pl);sp.insert(1,geom);sp.append(OxmlElement('a:effectLst'))
  sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();clean(sh)
 return dict(id=name,formula_width_pt=vw,formula_height_pt=vh,allocated_box=box,scale=1,outline_paths=len(paths),fits=vw<=w+.2 and vh<=h+.2)
