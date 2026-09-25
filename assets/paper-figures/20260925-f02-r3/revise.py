"""R3 directly edits copies of the native R2 PPTX; no source raster redraw."""
from pathlib import Path
import json,hashlib,copy,sys
import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree
import native_tools as nt

HERE=Path(__file__).resolve().parent;R=HERE.parent
SOURCE=R.parent/'final/history_unified_20260925/R2'
OUT=HERE/(sys.argv[1] if len(sys.argv)>1 else 'build01');OUT.mkdir(exist_ok=True);nt.OUT=OUT

def freeform_loop(sl,points):
 """One open native custom-geometry path, not a group of disjoint lines."""
 x=min(p[0] for p in points);y=min(p[1] for p in points)
 w=max(p[0] for p in points)-x;h=max(p[1] for p in points)-y
 sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.name='next_rollout_continuous'
 sp=sh._element.spPr
 for e in list(sp):
  if e.tag.endswith('prstGeom'):sp.remove(e)
 geom=OxmlElement('a:custGeom')
 for tag in ('avLst','gdLst','ahLst','cxnLst'):geom.append(OxmlElement('a:'+tag))
 rect=OxmlElement('a:rect')
 for k,v in dict(l='0',t='0',r='r',b='b').items():rect.set(k,v)
 geom.append(rect);pl=OxmlElement('a:pathLst');path=OxmlElement('a:path');path.set('w',str(Pt(w)));path.set('h',str(Pt(h)));path.set('fill','none');path.set('stroke','1')
 for i,(xx,yy) in enumerate(points):
  op=OxmlElement('a:moveTo' if i==0 else 'a:lnTo');pt=OxmlElement('a:pt');pt.set('x',str(Pt(xx-x)));pt.set('y',str(Pt(yy-y)));op.append(pt);path.append(op)
 pl.append(path);geom.append(pl);sp.insert(1,geom)
 sh.fill.background();sh.line.color.rgb=RGBColor.from_string(nt.GRAY);sh.line.width=Pt(.7)
 ln=sp.find('{'+nt.ANS+'}ln');end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');ln.append(end)
 nt.clean(sh);return sh

def build(n):
 src=SOURCE/f'Figure{n}_F02_R2.pptx';prs=Presentation(src);sl=prs.slides[0]
 assert prs.slide_width==Pt(396) and prs.slide_height==Pt(237.6)
 baseline={s.name:etree.tostring(s._element).decode() for s in sl.shapes}
 audit=json.loads((SOURCE/f'Figure{n}_formula_audit.json').read_text())
 font_changes=[]
 def raise_floor(shapes):
  for s in shapes:
   if s.shape_type==6:raise_floor(s.shapes)
   if s.has_text_frame:
    for p in s.text_frame.paragraphs:
     for run in p.runs:
      if run.font.size and run.font.size.pt<8.1:
       font_changes.append({'shape':s.name,'before_pt':run.font.size.pt,'after_pt':8.1});run.font.size=Pt(8.1)
 raise_floor(sl.shapes)
 if n==1:
  nt.text(sl,'penalty_schematic','(schematic)',10,32,67,10,8.1,nt.GRAY)
  nt.text(sl,'lagrangian_schematic','(schematic)',317,32,67,10,8.1,nt.GRAY)
  changed_math=set()
 else:
  # Lift only the lower update contents by 4pt and shorten its background.
  # This preserves the complete 35pt advantage allocation while making room
  # for a readable loop label below the panel and a continuous outer path.
  up={'sum','sum_sign','actor_step','actor_step_text','beta_update','lambda_out','advantage','pi_next'}
  prefixes=('beta_to_sum','sum_to_price','price_to_adv','adv_to_actor','actor_to_policy')
  for s in sl.shapes:
   if s.name in up or any(s.name.startswith(k) for k in prefixes):s.top-=Pt(4)
  nt.box(sl,'update',h=45)
  changed_math={'gae','beta_update','advantage'}
  for a in audit:
   if a['id'] in up:a['allocated_box'][1]-=4
   if a['id']=='gae':a['latex']=r'A_r,\ A_{c_k}'
   if a['id']=='beta_update':a['latex']=r'\beta_k\leftarrow[\beta_k+\alpha(\bar J_k-B_k)]_+'
   if a['id']=='advantage':a['latex']=r'\widetilde A=\dfrac{A_r-\sum_k\lambda_k A_{c_k}}{1+\sum_k\lambda_k}'
  nt.line(sl,'processed_to_sum',[(144,160),(137,160),(137,193)],nt.TEAL,.8,True)
  nt.line(sl,'gae_route',[(65,114),(82,114),(82,145),(300,145),(300,181),(282,181),(282,185)],nt.GRAY,.7,True)
  # Observed cost takes the internal gap; it no longer shares the outer left
  # margin with the return loop. It enters beta above the formula's right side.
  nt.line(sl,'cost_route',[(18,138),(18,153),(111,153),(111,187)],nt.ORANGE,.7,True)
  # Keep the existing white-backed data label above the regenerated GAE line.
  label=nt.get(sl,'gae_arrow_label');sl.shapes._spTree.remove(label._element);sl.shapes._spTree.append(label._element)
  nt.delete(sl,names=('next_label',),prefixes=('next_rollout',))
  loop=[(391,201),(394.5,201),(394.5,233),(1.2,233),(1.2,30.5),(27,30.5)]
  freeform_loop(sl,loop)
  nt.text(sl,'next_label','next rollout',169,221,64,10,8.1,nt.GRAY)
  (OUT/'loop_geometry.json').write_text(json.dumps({'path_count':1,'closed':False,'points_pt':loop,'stroke_pt':.7,'endpoint':'pi_t','source':'pi_t+1','outer_x_pt':[1.2,394.5],'panel_x_span_pt':[3,393],'bottom_path_y_pt':233,'panel_bottom_y_pt':219,'label_box_pt':[169,221,64,10],'label_bottom_margin_to_canvas_pt':6.6,'cost_vertical_x_pt':111,'left_route_separation_pt':109.8},indent=2))
 # Use R2's exact typesetting sizes. Unchanged outlines remain the existing
 # objects; new alpha/unhatted expressions alone are reimported at scale1.
 oldsource=HERE.parent/'revision_r2/revise.py'
 import ast
 tree=ast.parse(oldsource.read_text());eqfn=next(x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name=='equation_rows')
 env={'INK':nt.INK,'TEAL':nt.TEAL,'ORANGE':nt.ORANGE};exec(compile(ast.Module(body=[eqfn],type_ignores=[]),str(oldsource),'exec'),env)
 oldrows={row[0]:row for row in env['equation_rows'](n)}
 rows=[(a['id'],a['latex'],a['allocated_box'],oldrows[a['id']][3],oldrows[a['id']][4]) for a in audit]
 d=nt.typeset(rows,n);pdf=fitz.open(d/'equations.pdf');svgs=sorted(d.glob('eq-*.svg'),key=lambda p:int(p.stem.split('-')[-1]))
 assert len(svgs)==len(rows)==len(pdf)
 for i,(row,page,svg) in enumerate(zip(rows,pdf,svgs)):
  name,tex,b,size,c=row
  if name in changed_math:
   nt.delete(sl,names=(name,));a=nt.import_formula(sl,svg,name,b,c,page.rect);a['latex']=tex;audit[i]=a
  audit[i]['minimum_latex_pdf_font_pt']=min(s['size'] for bl in page.get_text('dict')['blocks'] for l in bl.get('lines',[]) for s in l['spans'] if s['text'].strip())
  audit[i]['allocated_box']=b
  assert page.rect.width<=b[2]+.2 and page.rect.height<=b[3]+.2,(name,page.rect,b)
 assert all(a['minimum_latex_pdf_font_pt']>=8 for a in audit)
 props=prs.core_properties;props.author='';props.last_modified_by='';props.comments='';props.title=f'SafeOT Figure {n} R3';props.subject='';props.keywords=''
 if sl.has_notes_slide:sl.notes_slide.notes_text_frame.text=''
 target=d/f'Figure{n}_F02_R3.pptx';prs.save(target)
 after={s.name:etree.tostring(s._element).decode() for s in sl.shapes}
 changes={'source_pptx':src.name,'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'identical_objects':[k for k in baseline if after.get(k)==baseline[k]],'modified_objects':[k for k in baseline if k in after and after[k]!=baseline[k]],'removed_objects':[k for k in baseline if k not in after],'added_objects':[k for k in after if k not in baseline],'font_floor_changes':font_changes,'math_reimported':sorted(changed_math),'pipeline':'R2 PPTX objects edited in place on a copy; 1:1 pdflatex/newtx outlines; no images'}
 (d/'object_changes.json').write_text(json.dumps(changes,indent=2));(d/'formula_audit.json').write_text(json.dumps(audit,indent=2))
 print(json.dumps({'figure':n,'output':str(target),'unchanged_native_objects':len(changes['identical_objects']),'modified_objects':changes['modified_objects'],'new_objects':changes['added_objects']},indent=2))

if __name__=='__main__':
 for n in (1,2):build(n)
