from pathlib import Path
import hashlib,json,shutil
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
import native_tools as nt
H=Path(__file__).resolve().parent;R=H.parent;B=R.parent/'final/history_unified_20260925/R4';O=H/'build02';O.mkdir(exist_ok=False)
for n in (1,2):
 src=B/f'Figure{n}_F02_R4.pptx';p=Presentation(src);sl=p.slides[0]
 before={s.name:etree.tostring(s._element,method='c14n') for s in sl.shapes}
 audit=json.loads((B/f'Figure{n}_formula_audit.json').read_text())
 if n==1:
  for name in ('leg_observed','leg_target','Fhat_legend','Fstar_legend'):nt.get(sl,name).left-=Pt(6)
  for a in audit:
   if a['id'] in ('Fhat_legend','Fstar_legend'):a['allocated_box'][0]-=6
  nt.text(sl,'upper_cost_label','high cost',177,29,36,11,8.1,nt.ORANGE)
  nt.line(sl,'upper_cost_tick',[(195,39.5),(195,45.5)],nt.ORANGE,.6)
  allowed={'leg_observed','leg_target','Fhat_legend','Fstar_legend'}
 else:
  # Match actual visible lower margins of panels1-3 to about2.6pt.
  for name in ('observed_J','observed_J_label'):nt.get(sl,name).top-=Pt(2.8)
  nt.get(sl,'edge_inputs').top+=Pt(13.3)
  nt.get(sl,'hard_label').top-=Pt(.4)
  for a in audit:
   if a['id']=='observed_J':a['allocated_box'][1]-=2.8
   if a['id']=='edge_inputs':a['allocated_box'][1]+=13.3
  route=nt.get(sl,'cost_route_0');route.top=Pt(134.2);route.height=Pt(18.8)
  label=nt.get(sl,'flow_balance');runs=[r for p in label.text_frame.paragraphs for r in p.runs];assert len(runs)==1
  runs[0].text='s.t. flow conservation + constraints'
  nt.text(sl,'upper_cost_label','high cost',274.5,63,31,11,8.1,nt.ORANGE)
  nt.line(sl,'upper_cost_tick',[(290,58),(290,63.8)],nt.ORANGE,.6)
  allowed={'observed_J','observed_J_label','edge_inputs','hard_label','cost_route_0','flow_balance'}
 # No formula contours or mathematics are regenerated.
 d=O/f'figure{n}';d.mkdir();p.save(d/f'Figure{n}_F02_R5.pptx')
 after={s.name:etree.tostring(s._element,method='c14n') for s in sl.shapes}
 modified=[k for k in before if k in after and after[k]!=before[k]];added=[k for k in after if k not in before];removed=[k for k in before if k not in after]
 assert set(modified)==allowed,(n,modified,allowed);assert set(added)=={'upper_cost_label','upper_cost_tick_0'} and not removed
 (d/'object_changes.json').write_text(json.dumps(dict(figure=n,source=src.name,source_sha256=hashlib.sha256(src.read_bytes()).hexdigest(),modified_objects=modified,added_objects=added,removed_objects=removed,identical_native_objects=sum(before[k]==after[k] for k in before),all_other_objects_identical=True,math_contours_unchanged=True,scope='Only21:12Z optional cost cue, legend placement, spacing, terminology repairs'),indent=2))
 (d/'formula_audit.json').write_text(json.dumps(audit,indent=2))
 shutil.copy2(B/f'Figure{n}_equations.tex',d/'equations.tex')
 print(n,modified,added)
