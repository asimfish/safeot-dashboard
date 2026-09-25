"""R2: edit the delivered PPTX objects; replace formula groups by LaTeX outlines.
The R1 files are inputs only. Every formula is one native editable vector group.
"""
from pathlib import Path
import copy,json,os,re,subprocess,sys,hashlib
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from fontTools.svgLib.path import parse_path
from fontTools.pens.recordingPen import RecordingPen
from fontTools.pens.transformPen import TransformPen
from fontTools.pens.boundsPen import BoundsPen
from fontTools.misc.transform import Transform
import fitz

HERE=Path(__file__).resolve().parent; R=HERE.parent
FINAL=R.parent/'final/history_unified_20260925'
OUT=HERE/'build04'; OUT.mkdir(exist_ok=True)
TEAL='087F8C'; GRAY='66747D'; INK='24343D'; LIGHT='A6B1B7'; ORANGE='BA5B20'; RED='AD3C3C'
PNS='http://schemas.openxmlformats.org/presentationml/2006/main'
ANS='http://schemas.openxmlformats.org/drawingml/2006/main'
MC='http://schemas.openxmlformats.org/markup-compatibility/2006'
ENV={**os.environ,'PATH':'/usr/bin:/bin:'+os.environ.get('PATH','')}

def delete(sl,names=(),prefixes=()):
 for s in list(sl.shapes):
  if s.name in names or any(s.name.startswith(x) for x in prefixes):s._element.getparent().remove(s._element)
def get(sl,name):return next(s for s in sl.shapes if s.name==name)
def box(sl,name,x=None,y=None,w=None,h=None):
 s=get(sl,name)
 for attr,v in zip(('left','top','width','height'),(x,y,w,h)):
  if v is not None:setattr(s,attr,Pt(v))
 return s
def clean(s):
 for r in s._element.findall('.//{'+ANS+'}effectRef'):r.set('idx','0')
def text(sl,name,t,x,y,w,h=11,size=8,color=INK,bold=False,align=PP_ALIGN.CENTER,fill=None):
 # LibreOffice rounds nominal 8 pt to 7.994 pt in this environment.
 if size==8:size=8.1
 delete(sl,[name]);s=sl.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));s.name=name
 tf=s.text_frame;tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 if fill:s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill)
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0)
  run=p.add_run();run.text=line;run.font.name='Liberation Serif';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 clean(s);return s
def line(sl,name,points,color=GRAY,width=.8,arrow=False,dash=False):
 delete(sl,prefixes=(name,))
 for i,(a,b) in enumerate(zip(points,points[1:])):
  s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Pt(v) for v in (*a,*b)]);s.name=name+'_'+str(i)
  s.line.color.rgb=RGBColor.from_string(color);s.line.width=Pt(width)
  ln=s._element.spPr.find('{'+ANS+'}ln')
  if dash:
   d=OxmlElement('a:prstDash');d.set('val','dash');ln.append(d)
  if arrow and i==len(points)-2:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
  clean(s)
def hard(sl,prefix,a,b):
 delete(sl,prefixes=(prefix+'_hard',prefix+'_x'))
 # Center-to-center edge painted behind the nodes, with a cross on the edge.
 line(sl,prefix+'_hard',[a,b],RED,.8,dash=True)
 cx=(a[0]+b[0])/2;cy=(a[1]+b[1])/2
 line(sl,prefix+'_x1',[(cx-2.4,cy-2.4),(cx+2.4,cy+2.4)],RED,1.2)
 line(sl,prefix+'_x2',[(cx-2.4,cy+2.4),(cx+2.4,cy-2.4)],RED,1.2)
 for s in list(sl.shapes):
  if s.name.startswith(prefix+'_node'):
   el=s._element;el.getparent().remove(el);sl.shapes._spTree.insert_element_before(el,'p:extLst')

def equation_rows(n):
 # Real TeX, explicit text spaces, fixed 8.2 TeX-point scripts (=8.169 PDF pt).
 if n==1:return [
 ('Fhat_legend',r'\widehat F',[111,29,15,14],11,INK),
 ('Fstar_legend',r'F^{*}',[147,29,17,14],11,INK),
 ('budget',r'\langle C_k,F\rangle\le b_k',[201,28,98,17],11,INK),
 ('raw_dual',r'\lambda_k^{*}',[167,118,23,17],11,TEAL),
 ('filter_chain',r'\text{clip + EMA}\;\to\;\bar{\lambda}^{*}_k',[143,137,109,17],10,TEAL),
 ('feedback_cost',r'\bar J_k-B_k',[94,169,43,13],10,ORANGE),
 ('price_sum',r'\lambda_k=\bar\lambda_k^{*}+\beta_k',[143,162,81,20],11,INK),
 ('priced_label',r'\widetilde A',[227,159,20,14],11,INK),
 ('pi_next',r'\pi_{t+1}',[312,164,30,17],11,INK),
 ('u_fixed_eq',r'c=0,\;\beta_k\ \text{fixed}',[7,219,80,15],8.5,INK),
 ('u_lag_eq',r'c=0,\;\beta_k\ \text{adaptive}',[88,219,90,15],8.5,INK),
 ('u_safe_eq',r'\text{finite}\ \varepsilon',[180,219,58,15],9,INK),
 ('u_lp_eq',r'\varepsilon\to0',[239,219,52,15],10,INK),
 ('u_active_eps',r'\varepsilon\to\infty',[294,215,95,13],9,INK),
 ]
 return [
 ('pi_t',r'\pi_t',[27,22,26,17],11,INK),
 ('gae',r'\widehat A_r,\ \widehat A_{c_k}',[9,102,64,17],11,INK),
 ('observed_J',r'\bar J_k',[9,121,19,17],10,ORANGE),
 ('edge_inputs',r'\widehat F,\ A,\ C_k',[93,108,62,18],11,INK),
 ('soft_eq',r'\langle C_k,F\rangle\le b_k',[170,34,77,17],11,INK),
 ('budget_scale',r'b_k=B_k/L',[175,52,67,16],10,INK),
 ('hard_eq',r'F_H=0',[182,112,51,16],11,INK),
 ('objective',r'\min_{F\ge0}\! -\langle A,F\rangle+\varepsilon\,\mathrm{KL}(F\Vert\widehat F)',[258,24,132,23],10,INK),
 ('before_F',r'\widehat F',[274,88,23,17],11,INK),
 ('after_F',r'F^{*}',[345,88,23,17],11,INK),
 ('filter_chain',r'\bar\lambda_k^{*}\ \longleftarrow\ \text{clip + EMA}\ \longleftarrow\ \lambda_k^{*}',[145,150,148,18],10,TEAL),
 ('beta_update',r'\beta_k\leftarrow[\beta_k+\eta(\bar J_k-B_k)]_+',[8,193,117,23],10,ORANGE),
 ('lambda_out',r'\lambda_k',[151,195,24,18],11,INK),
 ('advantage',r'\widetilde A=\dfrac{\widehat A_r-\sum_k\lambda_k\widehat A_{c_k}}{1+\sum_k\lambda_k}',[186,187,113,35],11,INK),
 ('pi_next',r'\pi_{t+1}',[363,195,29,18],11,INK),
 ]

def modify(n):
 src=FINAL/f'Figure{n}_F02_editable.pptx';prs=Presentation(src);sl=prs.slides[0]
 for el in list(sl.shapes._spTree):
  if el.tag=='{'+MC+'}AlternateContent':sl.shapes._spTree.remove(el)
 prs.slide_height=Pt(237.6)
 if n==1:
  for name in ('penalty','safeot','lagrangian'):box(sl,name,h=151)
  text(sl,'hard_event_label','hard event: excluded edge',139,82,116,11,8,RED)
  # A short legend key labels the unique red excluded edge without crossing flow arrows.
  line(sl,'hard_key',[(123,88),(135,88)],RED,.7,dash=True)
  line(sl,'hard_key_x1',[(127,86),(131,90)],RED,.9)
  line(sl,'hard_key_x2',[(127,90),(131,86)],RED,.9)
  text(sl,'graph_legend','nodes: clustered states · edges: transition flow',96,92,203,10,8)
  box(sl,'projection',y=104);box(sl,'projection_text',y=106)
  line(sl,'duals_down',[(195,120),(195,136)],TEAL,.8,True)
  text(sl,'pre_violation','Can price before violation',201,120,100,12,8,TEAL,True)
  delete(sl,['filter']);line(sl,'filter_to_price',[(233,153),(233,156),(182,156),(182,164)],TEAL,.8,True)
  box(sl,'price_rule',y=158,w=259,h=27)
  text(sl,'feedback_label','feedback',96,160,39,10,8,ORANGE)
  line(sl,'feedback_to_price',[(136,174),(142,174)],ORANGE,.8,True)
  line(sl,'price_to_actor',[(225,174),(249,174)],GRAY,.8,True)
  box(sl,'actor',x=251,y=165,w=43,h=17)
  text(sl,'actor_text','PPO/TRPO',252,168,41,11,8)
  line(sl,'actor_out',[(296,174),(309,174)],GRAY,.8,True)
  text(sl,'rule_title','One price rule',7,163,79,14,9,TEAL,True)
  text(sl,'exact_title','Exact settings',8,186,168,12,8,GRAY,True)
  text(sl,'limits_title','Conditional limits',248,186,141,12,8,GRAY,True)
  box(sl,'unification',y=200,h=35)
  line(sl,'exact_bracket',[(10,198),(10,200),(175,200),(175,198)],GRAY,.6)
  line(sl,'limit_bracket',[(248,198),(248,200),(388,200),(388,198)],GRAY,.6)
  for name,title,x,w in [('fixed','Fixed penalty',7,80),('lag','Lagrangian',88,90),('ours','SafeOT',180,58),('lp','LP',239,52),('active','Active-set',294,95)]:
   text(sl,'u_'+name,title,x,203,w,12,8.5,TEAL if name=='ours' else INK,True)
  text(sl,'u_active_eq','saturated clip',294,225,95,10,8)
 else:
  for name in ('rollout','inputs','budgets','projection'):box(sl,name,h=136)
  delete(sl,['rollout_label','observed'])
  text(sl,'GAE_label','Rollout + GAE',8,88,67,12,8,INK,True)
  text(sl,'observed_J_label','observed\nepisode cost',29,120,47,19,8,ORANGE,False,PP_ALIGN.LEFT)
  # Add room for b_k = B_k/L without shrinking the exclusion graph.
  for s in list(sl.shapes):
   if s.name.startswith('excluded_graph'):s.top+=Pt(13)
  hard(sl,'excluded_graph',(195.58,77),(220.42,106))
  hard(sl,'after',(347.88,58),(368.12,82))
  text(sl,'hard_label','Hard exclusions',173,127,71,10,8)
  text(sl,'flow_balance','s.t. flow balance + constraints',260,43,130,11,7.5)
  text(sl,'target_only','Target flow is not executed',259,113,130,12,8,TEAL,True)
  box(sl,'update',y=174,h=54)
  text(sl,'update_title','5. Priced update',8,177,91,12,9,INK,True,PP_ALIGN.LEFT)
  box(sl,'sum',x=131,y=199,w=12,h=12)
  text(sl,'sum_sign','+',132,200,10,10,8.5)
  line(sl,'beta_to_sum',[(124,205),(129,205)],ORANGE,.8,True)
  line(sl,'sum_to_price',[(145,205),(150,205)],GRAY,.8,True)
  line(sl,'price_to_adv',[(177,205),(184,205)],GRAY,.8,True)
  line(sl,'adv_to_actor',[(301,205),(311,205)],GRAY,.8,True)
  box(sl,'actor_step',x=313,y=195,w=43,h=22)
  text(sl,'actor_step_text','PPO/TRPO',314,200,41,12,8)
  line(sl,'actor_to_policy',[(357,205),(362,205)],GRAY,.8,True)
  # Gray advantages travel ABOVE the entire graph-price filter and descend to A-tilde.
  line(sl,'gae_route',[(65,114),(82,114),(82,145),(300,145),(300,185),(282,185),(282,189)],GRAY,.7,True)
  text(sl,'gae_arrow_label','advantages',91,140,48,10,8,GRAY,False,fill='FFFFFF')
  delete(sl,prefixes=('dual_in','processed_to_sum','raw_to_clip','clip_to_processed'))
  box(sl,'filter_box',x=144,y=151,w=151,h=18)
  delete(sl,['filter_text'])
  line(sl,'dual_in',[(326,140),(326,160),(297,160)],TEAL,.8,True)
  text(sl,'dual_label','budget duals',332,152,58,12,8,TEAL)
  line(sl,'processed_to_sum',[(144,160),(137,160),(137,197)],TEAL,.8,True)
  # This orange line starts at the explicitly labeled observed-cost output, not GAE.
  line(sl,'cost_route',[(18,138),(18,163),(7,163),(7,195)],ORANGE,.7,True)
  line(sl,'next_rollout',[(379,217),(379,234),(1,234),(1,64),(3,64)],GRAY,.7,True)
  text(sl,'next_label','next rollout',172,226,58,10,8,GRAY,False,fill='FFFFFF')
 for s in sl.shapes:clean(s)
 props=prs.core_properties;props.author='';props.last_modified_by='';props.comments='';props.title=f'SafeOT Figure {n} R2'
 return prs,equation_rows(n)

def typeset(rows,n):
 d=OUT/f'figure{n}';d.mkdir(exist_ok=True)
 header=r'''\documentclass{article}
\usepackage{newtxtext,newtxmath}
\pagestyle{empty}
'''
 for s in sorted(set(r[3] for r in rows)):header+=rf'\DeclareMathSizes{{{s}}}{{{s}}}{{8.2}}{{8.2}}'+'\n'
 header+='\\begin{document}\n'
 for name,tex,b,size,color in rows:
  header+=rf'''\begingroup
\fontsize{{{size}}}{{14}}\selectfont
\setbox0\hbox{{$ {tex} $}}
\pdfpagewidth=\dimexpr\wd0+2pt\relax
\pdfpageheight=\dimexpr\ht0+\dp0+2pt\relax
\pdfhorigin=0pt \pdfvorigin=0pt
\shipout\vbox{{\offinterlineskip\kern1pt\hbox{{\kern1pt\box0\kern1pt}}\kern1pt}}
\endgroup
'''
 header+='\\end{document}\n';(d/'equations.tex').write_text(header)
 p=subprocess.run(['/usr/bin/pdflatex','-no-shell-escape','-interaction=batchmode','-halt-on-error','equations.tex'],cwd=d,env=ENV,capture_output=True,text=True)
 (d/'tex_build.log').write_text(p.stdout+p.stderr);assert p.returncode==0,(d/'equations.log').read_text()[-3000:]
 p=subprocess.run(['/usr/bin/dvisvgm','--pdf','--no-fonts','--exact','--page=1-','--output=eq-%p.svg','equations.pdf'],cwd=d,capture_output=True,text=True)
 (d/'svg_build.log').write_text(p.stdout+p.stderr);assert p.returncode==0,p.stderr
 return d

def import_formula(sl,svg,name,box,color,page_rect):
 root=etree.parse(str(svg)).getroot()
 # dvisvgm 2.13 accumulates the max media box across pages. Each page's
 # original PDF rectangle gives the correct origin; glyph coordinates are unchanged.
 vb=[0,-page_rect.height,page_rect.width,page_rect.height];vw,vh=vb[2:]
 x,y,w,h=box;ox=x+(w-vw)/2-vb[0];oy=y+(h-vh)/2-vb[1]
 if vw>w+.2 or vh>h+.2:print('FORMULA_OVERFLOW',name,vw,vh,box)
 g=sl.shapes.add_group_shape();g.name=name
 paths=[]
 def walk(node,tr=Transform()):
  val=node.get('transform','')
  if val:
   nums=list(map(float,re.findall(r'[-+]?(?:\d*\.)?\d+(?:[eE][-+]?\d+)?',val)));assert val.startswith('matrix('),val
   tr=tr.transform(Transform(*nums))
  if node.tag.endswith('}path'):
   rec=RecordingPen();parse_path(node.get('d'),TransformPen(rec,tr.translate(0,0)))
   pts=BoundsPen(None);rec.replay(pts)
   if pts.bounds:paths.append((rec,pts.bounds))
  for ch in node:walk(ch,tr)
 walk(root)
 for idx,(rec,(xmin,ymin,xmax,ymax)) in enumerate(paths):
  ew=max(.02,xmax-xmin);eh=max(.02,ymax-ymin)
  sh=g.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(ox+xmin),Pt(oy+ymin),Pt(ew),Pt(eh));sh.name=f'{name}_outline_{idx}'
  sp=sh._element.spPr
  for z in list(sp):
   if z.tag.endswith('prstGeom'):sp.remove(z)
  geom=OxmlElement('a:custGeom')
  for tag in ('avLst','gdLst','ahLst','cxnLst'):geom.append(OxmlElement('a:'+tag))
  rect=OxmlElement('a:rect')
  for k,v in dict(l='0',t='0',r='r',b='b').items():rect.set(k,v)
  geom.append(rect);pl=OxmlElement('a:pathLst');path=OxmlElement('a:path');path.set('w',str(round(ew*12700)));path.set('h',str(round(eh*12700)))
  for op,points in rec.value:
   if op in ('closePath','endPath'):
    if op=='closePath':path.append(OxmlElement('a:close'))
    continue
   tag={'moveTo':'moveTo','lineTo':'lnTo','curveTo':'cubicBezTo','qCurveTo':'quadBezTo'}[op]
   node=OxmlElement('a:'+tag)
   for xx,yy in points:
    pt=OxmlElement('a:pt');pt.set('x',str(round((xx-xmin)*12700)));pt.set('y',str(round((yy-ymin)*12700)));node.append(pt)
   path.append(node)
  pl.append(path);geom.append(pl);sp.insert(1,geom);sp.append(OxmlElement('a:effectLst'))
  sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();clean(sh)
 return dict(id=name,formula_width_pt=vw,formula_height_pt=vh,allocated_box=box,scale=1,outline_paths=len(paths),fits=vw<=w+.2 and vh<=h+.2)

def build(n):
 prs,rows=modify(n);d=typeset(rows,n);pdf=fitz.open(d/'equations.pdf');audit=[]
 for i,(name,tex,b,size,col) in enumerate(rows):
  a=import_formula(prs.slides[0],sorted(d.glob('eq-*.svg'))[i],name,b,col,pdf[i].rect)
  a.update(latex=tex,minimum_latex_pdf_font_pt=min(s['size'] for bl in pdf[i].get_text('dict')['blocks'] for ln in bl.get('lines',[]) for s in ln['spans'] if s['text'].strip()))
  audit.append(a)
 target=d/f'Figure{n}_F02_R2.pptx';prs.save(target)
 (d/'formula_audit.json').write_text(json.dumps(audit,indent=2))
 (d/'provenance.json').write_text(json.dumps({'source_pptx':str(FINAL/f'Figure{n}_F02_editable.pptx'),'source_sha256':hashlib.sha256((FINAL/f'Figure{n}_F02_editable.pptx').read_bytes()).hexdigest(),'formulas':'Actual LaTeX newtx outlines, grouped native freeforms at scale 1; no alternate rendering branch','powerpoint_render_verified':False},indent=2))
 print(target)

if __name__=='__main__':
 for n in ([int(sys.argv[1])] if len(sys.argv)>1 else [1,2]):build(n)
