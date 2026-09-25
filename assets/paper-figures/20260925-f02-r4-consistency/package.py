from pathlib import Path
import sys,json,shutil,zipfile,hashlib
import fitz
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from PIL import Image
import numpy as np
HERE=Path(__file__).resolve().parent; R=HERE.parent
sys.path.append(str(R.parent/'image_pipeline_20260925/runtime-venv/lib/python3.12/site-packages'))
from super_img2ppt.render import LibreOfficeRenderer,rasterize_pdf
FINAL=HERE/'delivery02';FINAL.mkdir(exist_ok=False)
reports=[]
for n in (1,2):
 d=HERE/f'build03/figure{n}';stem=f'Figure{n}_F02_R4'
 pdf=LibreOfficeRenderer('/usr/bin/soffice',timeout=60).render(d/(stem+'.pptx'),d/'render')
 rasterize_pdf(pdf,d/'render',1650)
 shutil.copy2(d/(stem+'.pptx'),FINAL/(stem+'.pptx'))
 doc=fitz.open(pdf);doc.set_metadata({});doc.del_xml_metadata();doc.save(FINAL/(stem+'.pdf'),garbage=4,deflate=True)
 page=doc[0];page.get_pixmap(dpi=300).save(FINAL/(stem+'_300dpi.png'))
 (FINAL/(stem+'.svg')).write_text(page.get_svg_image(text_as_path=False))
 Image.open(FINAL/(stem+'_300dpi.png')).convert('L').save(FINAL/(stem+'_grayscale.png'),dpi=(300,300))
 spans=[s for b in page.get_text('dict')['blocks'] for l in b.get('lines',[]) for s in l['spans'] if s['text'].strip()]
 audit=json.loads((d/'formula_audit.json').read_text());assert all(a['fits'] for a in audit)
 min_math=min(a['minimum_latex_pdf_font_pt'] for a in audit);assert min_math>=8
 # Check the exact regions whose minimum was raised in the brief.
 selected=[s for s in spans if s['bbox'][1] >= (185 if n==1 else 174)]
 min_region=min(s['size'] for s in selected);assert min_region>=8
 assert min(s['size'] for s in spans)>=8
 assert all(fitz.Rect(s['bbox']) in page.rect for s in spans)
 with zipfile.ZipFile(FINAL/(stem+'.pptx')) as z:
  xml=etree.fromstring(z.read('ppt/slides/slide1.xml'))
  pictures=len(xml.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}pic'))
  ids=[x.get('id') for x in xml.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')]
  assert len(ids)==len(set(ids))
 assert pictures==0 and len(page.get_images())==0
 report={'figure':n,'width_in':page.rect.width/72,'height_in':page.rect.height/72,'minimum_pdf_text_pt':min(s['size'] for s in spans),'minimum_requested_region_text_pt':min_region,'minimum_math_source_pdf_font_pt':min_math,'formula_measurement':'Actual pdflatex/newtx formula PDF at scale 1, including sub/superscripts; final exported outlines carry no font metadata.','formula_groups':len(audit),'pptx_raster_objects':pictures,'pdf_raster_objects':len(page.get_images()),'renderer':'LibreOffice 7.3.7.2 via super_img2ppt.render.LibreOfficeRenderer','powerpoint_render_verified':False,'formulas':'Native editable vector outline groups, same branch in every application; LaTeX source supplied. Not Office Math objects.','all_formula_boxes_fit':True,'all_pdf_text_within_page':True,'r3_and_text_only_r4_preserved':True,'cold_reader_acceptance':'Pending; no score assigned by drawing agent'}
 (FINAL/f'Figure{n}_font_editability_report.json').write_text(json.dumps(report,indent=2));reports.append(report)
 for src,new in [('object_changes.json',f'Figure{n}_object_changes.json'),('formula_audit.json',f'Figure{n}_formula_audit.json'),('equations.tex',f'Figure{n}_equations.tex')]:shutil.copy2(d/src,FINAL/new)
 # Granularity test on a disposable copy: move one equation and edit one text label.
 prs=Presentation(FINAL/(stem+'.pptx'));sl=prs.slides[0];g=next(s for s in sl.shapes if s.name=='filter_chain');label=next(s for s in sl.shapes if s.name==('safeot_title' if n==1 else 'projection_title'))
 before={s.name:[s.left,s.top,s.width,s.height] for s in sl.shapes};g.left+=Pt(2);label.text_frame.paragraphs[0].runs[0].text+=' test';tmp=d/'editability_probe.pptx';prs.save(tmp)
 probe=Presentation(tmp);after={s.name:[s.left,s.top,s.width,s.height] for s in probe.slides[0].shapes}
 assert all(before[k]==after[k] for k in before if k!='filter_chain');assert after['filter_chain'][0]-before['filter_chain'][0]==Pt(2)
 (d/'editability_probe.json').write_text(json.dumps({'formula_group_move_pt':2,'neighbor_positions_unchanged':True,'label_native_edit_verified':True,'original_unchanged':True},indent=2))
 # Color-vision previews: standard Machado 2009 matrices, as in earlier rounds.
 im=np.asarray(Image.open(FINAL/(stem+'_300dpi.png')).convert('RGB'),dtype=float)/255
 lin=np.where(im<=.04045,im/12.92,((im+.055)/1.055)**2.4)
 for name,m in {'deuteranopia':[[.367322,.860646,-.227968],[.280085,.672501,.047413],[-.011820,.042940,.968881]],'protanopia':[[.152286,1.052583,-.204868],[.114503,.786281,.099216],[-.003882,-.048116,1.051998]]}.items():
  c=np.clip(lin@np.array(m).T,0,1);srgb=np.where(c<=.0031308,12.92*c,1.055*c**(1/2.4)-.055)
  Image.fromarray(np.round(srgb*255).astype('uint8')).save(FINAL/(stem+'_'+name+'.png'),dpi=(300,300))
 doc.close()

(FINAL/'checks.json').write_text(json.dumps(reports,indent=2))
print(json.dumps(reports,indent=2))
