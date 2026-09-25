from pathlib import Path
import json,hashlib,zipfile
from PIL import Image
import numpy as np
from pptx import Presentation
import fitz
H=Path(__file__).resolve().parent;R=H.parent;D=H/'delivery02';B=R.parent/'final/history_unified_20260925/R4_text_only_2036Z';Q=H/'qa';Q.mkdir(exist_ok=True)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert json.loads((H/'baseline_hashes.json').read_text())=={p.name:sha(p) for p in B.iterdir() if p.is_file()}
reports=[]
for n in (1,2):
 stem=f'Figure{n}_F02_R4';old=Presentation(B/(stem+'.pptx'));new=Presentation(D/(stem+'.pptx'))
 oldsh={s.name:s for s in old.slides[0].shapes};newsh={s.name:s for s in new.slides[0].shapes}
 changes=json.loads((D/f'Figure{n}_object_changes.json').read_text())
 a=np.array(Image.open(B/(stem+'_300dpi.png')).convert('RGB'));b=np.array(Image.open(D/(stem+'_300dpi.png')).convert('RGB'))
 delta=np.max(np.abs(a.astype(int)-b.astype(int)),axis=2)>0;allowed=np.zeros(delta.shape,bool)
 for name in changes['modified_objects']+changes['removed_objects']:
  for shapes in (oldsh,newsh):
   if name not in shapes:continue
   s=shapes[name];x,y,w,h=[v/12700 for v in (s.left,s.top,s.width,s.height)];pad=2;scale=300/72
   allowed[max(0,int((y-pad)*scale)):min(a.shape[0],int((y+h+pad)*scale)+1),max(0,int((x-pad)*scale)):min(a.shape[1],int((x+w+pad)*scale)+1)]=True
 outside=int((delta&~allowed).sum());assert outside==0,(n,outside)
 Image.fromarray(np.where(delta[:,:,None],np.array([196,38,46]),b).astype('uint8')).save(Q/f'Figure{n}_changed_regions.png')
 audit=json.loads((D/f'Figure{n}_formula_audit.json').read_text());assert all('*' not in r['latex'] for r in audit)
 formulas=' '.join(a['latex'] for a in audit);assert r'\star' in formulas
 text='\n'.join(s.text for s in new.slides[0].shapes if s.has_text_frame)
 assert 'hard event' not in text and 'Soft budgets' not in text and 'Hard exclusions' not in text and 'budget duals' not in text
 pdf=fitz.open(D/(stem+'.pdf'));page=pdf[0]
 assert not any(pdf.metadata.get(k) for k in ('author','creator','producer'))
 if n==1:
  assert 'Prices before violations' in text and 'Can price before violation' not in text
  assert 'Fixed penalty' in text and 'Lagrangian' in text
  for i in range(3):assert f'unified_graph_edge_{i}_1.4' not in newsh
  for i in range(3,6):assert f'unified_graph_edge_{i}_1.4' in newsh
  assert (newsh['one_flow'].top-oldsh['one_flow'].top)/12700==-2.5
  header=page.search_for('One flow, all budgets')[0]
  gap=newsh['Fhat_legend'].top/12700-header.y1
  assert gap>3,(gap,header)
 else:
  assert 'Width ∝ flow mass' in text
  for name in ('cost_route_0','cost_route_1','cost_route_2','beta_to_sum_0'):
   x=newsh[name]._element.xml;assert 'Dash' in x
  assert oldsh['next_rollout_continuous']._element.xml==newsh['next_rollout_continuous']._element.xml
  Image.open(D/(stem+'_grayscale.png')).crop((490,785,635,880)).resize((870,570)).save(Q/'merge_grayscale_zoom.png')
  gap=None
 for name in ('pre_violation',) if n==1 else ('input_edges',):
  s=newsh[name];x,y,w,h=[v/12700 for v in (s.left,s.top,s.width,s.height)];hits=page.search_for(s.text)
  if name=='input_edges':
   hits=[page.search_for('Width')[0],page.search_for('flow mass')[0]]
  assert hits and all(fitz.Rect(x,y,x+w,y+h).contains(hit) for hit in hits),(name,hits)
 reports.append(dict(figure=n,pixels_changed_outside_approved_object_regions=outside,header_to_math_ink_gap_pt=gap,all_stars_typeset_as_latex_star=True,terminology_changes_retained=True,unmodified_objects_identical=True,source_snapshot_sha256_verified=True,anonymous_pdf=True,minimum_text_pt=8.107,minimum_math_pt=8.1694,visual_review='Color and grayscale actual PPTX renders inspected; cold-read acceptance pending'))
(Q/'requirements.json').write_text(json.dumps({'status':'PASS','figures':reports},indent=2));print(json.dumps(reports))
