from pathlib import Path

import copy,json,os,re,subprocess,sys,hashlib

from lxml import etree

from pptx import Presentation

from pptx.util import Pt

from pptx.dml.color import RGBColor

from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR

from pptx.enum.text import PP_ALIGN,MSO_ANCHOR

from pptx.oxml.xmlchemy import OxmlElement

from fontTools.svgLib.path import parse_path

from fontTools.pens.recordingPen import RecordingPen

from fontTools.pens.transformPen import TransformPen

from fontTools.pens.boundsPen import BoundsPen

from fontTools.misc.transform import Transform

import fitz

HERE=Path(__file__).resolve().parent

OUT=HERE/"build01"; OUT.mkdir(exist_ok=True)

ANS="http://schemas.openxmlformats.org/drawingml/2006/main"

TEAL="087F8C"; GRAY="66747D"; INK="24343D"; ORANGE="BA5B20"

ENV={**os.environ,"PATH":"/usr/bin:/bin:"+os.environ.get("PATH","")}

def delete(sl,names=(),prefixes=()):
 for s in list(sl.shapes):
  if s.name in names or any(s.name.startswith(x) for x in prefixes):s._element.getparent().remove(s._element)

def get(sl,name):return next(s for s in sl.shapes if s.name==name)

def box(sl,name,x=None,y=None,w=None,h=None):
 s=get(sl,name)
 for attr,v in zip(('left','top','width','height'),(x,y,w,h)):
  if v is not None:setattr(s,attr,Pt(v))
 return s

def clean(s):
 for r in s._element.findall('.//{'+ANS+'}effectRef'):r.set('idx','0')

def text(sl,name,t,x,y,w,h=11,size=8,color=INK,bold=False,align=PP_ALIGN.CENTER,fill=None):
 # LibreOffice rounds nominal 8 pt to 7.994 pt in this environment.
 if size==8:size=8.1
 delete(sl,[name]);s=sl.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));s.name=name
 tf=s.text_frame;tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 if fill:s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill)
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0)
  run=p.add_run();run.text=line;run.font.name='Liberation Serif';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 clean(s);return s

def line(sl,name,points,color=GRAY,width=.8,arrow=False,dash=False):
 delete(sl,prefixes=(name,))
 for i,(a,b) in enumerate(zip(points,points[1:])):
  s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Pt(v) for v in (*a,*b)]);s.name=name+'_'+str(i)
  s.line.color.rgb=RGBColor.from_string(color);s.line.width=Pt(width)
  ln=s._element.spPr.find('{'+ANS+'}ln')
  if dash:
   d=OxmlElement('a:prstDash');d.set('val','dash');ln.append(d)
  if arrow and i==len(points)-2:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
  clean(s)

def typeset(rows,n):
 d=OUT/f'figure{n}';d.mkdir(exist_ok=True)
 header=r'''\documentclass{article}
\usepackage{newtxtext,newtxmath}
\pagestyle{empty}
'''
 for s in sorted(set(r[3] for r in rows)):header+=rf'\DeclareMathSizes{{{s}}}{{{s}}}{{8.2}}{{8.2}}'+'\n'
 header+='\\begin{document}\n'
 for name,tex,b,size,color in rows:
  header+=rf'''\begingroup
\fontsize{{{size}}}{{14}}\selectfont
\setbox0\hbox{{$ {tex} $}}
\pdfpagewidth=\dimexpr\wd0+2pt\relax
\pdfpageheight=\dimexpr\ht0+\dp0+2pt\relax
\pdfhorigin=0pt \pdfvorigin=0pt
\shipout\vbox{{\offinterlineskip\kern1pt\hbox{{\kern1pt\box0\kern1pt}}\kern1pt}}
\endgroup
'''
 header+='\\end{document}\n';(d/'equations.tex').write_text(header)
 p=subprocess.run(['/usr/bin/pdflatex','-no-shell-escape','-interaction=batchmode','-halt-on-error','equations.tex'],cwd=d,env=ENV,capture_output=True,text=True)
 (d/'tex_build.log').write_text(p.stdout+p.stderr);assert p.returncode==0,(d/'equations.log').read_text()[-3000:]
 p=subprocess.run(['/usr/bin/dvisvgm','--pdf','--no-fonts','--exact','--page=1-','--output=eq-%p.svg','equations.pdf'],cwd=d,capture_output=True,text=True)
 (d/'svg_build.log').write_text(p.stdout+p.stderr);assert p.returncode==0,p.stderr
 return d

def import_formula(sl,svg,name,box,color,page_rect):
 root=etree.parse(str(svg)).getroot()
 # dvisvgm 2.13 accumulates the max media box across pages. Each page's
 # original PDF rectangle gives the correct origin; glyph coordinates are unchanged.
 vb=[0,-page_rect.height,page_rect.width,page_rect.height];vw,vh=vb[2:]
 x,y,w,h=box;ox=x+(w-vw)/2-vb[0];oy=y+(h-vh)/2-vb[1]
 if vw>w+.2 or vh>h+.2:print('FORMULA_OVERFLOW',name,vw,vh,box)
 g=sl.shapes.add_group_shape();g.name=name
 paths=[]
 def walk(node,tr=Transform()):
  val=node.get('transform','')
  if val:
   nums=list(map(float,re.findall(r'[-+]?(?:\d*\.)?\d+(?:[eE][-+]?\d+)?',val)));assert val.startswith('matrix('),val
   tr=tr.transform(Transform(*nums))
  if node.tag.endswith('}path'):
   rec=RecordingPen();parse_path(node.get('d'),TransformPen(rec,tr.translate(0,0)))
   pts=BoundsPen(None);rec.replay(pts)
   if pts.bounds:paths.append((rec,pts.bounds))
  for ch in node:walk(ch,tr)
 walk(root)
 for idx,(rec,(xmin,ymin,xmax,ymax)) in enumerate(paths):
  ew=max(.02,xmax-xmin);eh=max(.02,ymax-ymin)
  sh=g.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(ox+xmin),Pt(oy+ymin),Pt(ew),Pt(eh));sh.name=f'{name}_outline_{idx}'
  sp=sh._element.spPr
  for z in list(sp):
   if z.tag.endswith('prstGeom'):sp.remove(z)
  geom=OxmlElement('a:custGeom')
  for tag in ('avLst','gdLst','ahLst','cxnLst'):geom.append(OxmlElement('a:'+tag))
  rect=OxmlElement('a:rect')
  for k,v in dict(l='0',t='0',r='r',b='b').items():rect.set(k,v)
  geom.append(rect);pl=OxmlElement('a:pathLst');path=OxmlElement('a:path');path.set('w',str(round(ew*12700)));path.set('h',str(round(eh*12700)))
  for op,points in rec.value:
   if op in ('closePath','endPath'):
    if op=='closePath':path.append(OxmlElement('a:close'))
    continue
   tag={'moveTo':'moveTo','lineTo':'lnTo','curveTo':'cubicBezTo','qCurveTo':'quadBezTo'}[op]
   node=OxmlElement('a:'+tag)
   for xx,yy in points:
    pt=OxmlElement('a:pt');pt.set('x',str(round((xx-xmin)*12700)));pt.set('y',str(round((yy-ymin)*12700)));node.append(pt)
   path.append(node)
  pl.append(path);geom.append(pl);sp.insert(1,geom);sp.append(OxmlElement('a:effectLst'))
  sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();clean(sh)
 return dict(id=name,formula_width_pt=vw,formula_height_pt=vh,allocated_box=box,scale=1,outline_paths=len(paths),fits=vw<=w+.2 and vh<=h+.2)
