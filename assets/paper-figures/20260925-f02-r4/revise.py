"""R4 changes ONLY approved text nodes in immutable R3 PPTX copies."""
from pathlib import Path
import copy,json,hashlib,re,sys,shutil,zipfile
from lxml import etree
from pptx import Presentation
import fitz,numpy as np
from PIL import Image

HERE=Path(__file__).resolve().parent;R=HERE.parent
SOURCE=R.parent/'final/history_unified_20260925/R3'
FINAL=R.parent/'final/history_unified_20260925/R4';FINAL.mkdir(exist_ok=False)
sys.path.append(str(R.parent/'image_pipeline_20260925/runtime-venv/lib/python3.12/site-packages'))
from super_img2ppt.render import LibreOfficeRenderer
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
baseline_hashes={p.name:sha(p) for p in SOURCE.iterdir() if p.is_file()}
replacements={1:{'hard event: excluded edge':'hard constraint: excluded edge'},2:{'Soft budgets':'Soft constraints','Hard exclusions':'Hard constraints','budget duals':'safety prices'}}
checks=[]
for n in (1,2):
 src=SOURCE/f'Figure{n}_F02_R3.pptx';dst=FINAL/f'Figure{n}_F02_R4.pptx';stem=dst.stem
 with zipfile.ZipFile(src) as z:
  raw=z.read('ppt/slides/slide1.xml');root=etree.fromstring(raw);changes=[]
  for node in root.findall('.//a:t',NS):
   old=node.text or '';new=replacements[n].get(old,old)
   new=re.sub(r'\b(?:filtered|filter)\b','smoothed',new,flags=re.I)
   if old==new:continue
   sp=node
   while sp.tag!='{'+NS['p']+'}sp':sp=sp.getparent()
   name=sp.find('p:nvSpPr/p:cNvPr',NS).get('name')
   changes.append({'shape':name,'before':old,'after':new});node.text=new
  assert len(changes)==len(replacements[n]),changes
  newraw=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
  with zipfile.ZipFile(dst,'w') as out:
   for info in z.infolist():out.writestr(info,newraw if info.filename=='ppt/slides/slide1.xml' else z.read(info.filename))
 # Package and slide proof: restore approved strings in memory and compare XML.
 with zipfile.ZipFile(src) as a,zipfile.ZipFile(dst) as b:
  assert a.namelist()==b.namelist()
  altered=[f for f in a.namelist() if a.read(f)!=b.read(f)];assert altered==['ppt/slides/slide1.xml']
  oldxml=etree.fromstring(a.read(altered[0]));newxml=etree.fromstring(b.read(altered[0]))
  for x,y in zip(oldxml.findall('.//a:t',NS),newxml.findall('.//a:t',NS)):
   if x.text!=y.text:y.text=x.text
  assert etree.tostring(oldxml,method='c14n')==etree.tostring(newxml,method='c14n')
 prs=Presentation(dst);sl=prs.slides[0]
 boxes={s.name:[v/12700 for v in (s.left,s.top,s.width,s.height)] for s in sl.shapes}
 if n==1:
  assert next(s for s in sl.shapes if s.name=='penalty_title').text=='Fixed penalty'
  assert next(s for s in sl.shapes if s.name=='lagrangian_title').text=='Lagrangian'
 for c in changes:c['unchanged_box_pt']=boxes[c['shape']]
 build=HERE/f'figure{n}';build.mkdir(exist_ok=False)
 pdf=LibreOfficeRenderer('/usr/bin/soffice',timeout=60).render(dst,build/'render')
 doc=fitz.open(pdf);doc.set_metadata({});doc.del_xml_metadata();doc.save(FINAL/(stem+'.pdf'),garbage=4,deflate=True)
 page=doc[0];page.get_pixmap(dpi=300).save(FINAL/(stem+'_300dpi.png'));(FINAL/(stem+'.svg')).write_text(page.get_svg_image(text_as_path=False))
 im=Image.open(FINAL/(stem+'_300dpi.png')).convert('RGB');im.convert('L').save(FINAL/(stem+'_grayscale.png'),dpi=(300,300))
 spans=[s for b in page.get_text('dict')['blocks'] for l in b.get('lines',[]) for s in l['spans'] if s['text'].strip()]
 assert min(s['size'] for s in spans)>=8
 assert all(page.rect.contains(fitz.Rect(s['bbox'])) for s in spans)
 for c in changes:
  hits=page.search_for(c['after']);assert len(hits)==1,(c,hits)
  x,y,w,h=c['unchanged_box_pt'];rect=fitz.Rect(x,y,x+w,y+h);assert rect.contains(hits[0]),(c,rect,hits[0])
 assert not page.get_images()
 # Pixel differences must be confined to the changed text boxes (2px tolerance).
 a=np.asarray(Image.open(SOURCE/f'Figure{n}_F02_R3_300dpi.png').convert('RGB'));b=np.asarray(im);assert a.shape==b.shape
 delta=np.max(np.abs(a.astype(int)-b.astype(int)),axis=2)>0;allowed=np.zeros(delta.shape,dtype=bool)
 for c in changes:
  x,y,w,h=c['unchanged_box_pt'];scale=300/72
  xa=max(0,int(x*scale)-2);ya=max(0,int(y*scale)-2);xb=min(a.shape[1],int((x+w)*scale)+3);yb=min(a.shape[0],int((y+h)*scale)+3);allowed[ya:yb,xa:xb]=True
 outside=int(np.sum(delta&~allowed));assert outside==0,{'outside_approved_text_boxes':outside}
 mask=Image.fromarray(np.where(delta[:,:,None],np.array([192,40,45]),b).astype('uint8'));mask.save(build/'text_only_diff.png')
 for name in ['equations.tex','formula_audit.json']:
  shutil.copy2(SOURCE/f'Figure{n}_{name}',FINAL/f'Figure{n}_{name}')
 report=json.loads((SOURCE/f'Figure{n}_font_editability_report.json').read_text())
 report.update(minimum_pdf_text_pt=min(s['size'] for s in spans),source_r3_sha256=sha(src),r3_preserved=True,text_only_slide_xml_proof=True,pixels_changed_outside_approved_text_boxes=0,cold_reader_acceptance='Pending; text-only terminology revision')
 (FINAL/f'Figure{n}_font_editability_report.json').write_text(json.dumps(report,indent=2))
 proof=dict(figure=n,changes=changes,package_entries_changed=altered,all_other_package_entries_byte_identical=True,slide_XML_identical_except_approved_text=True,geometry_styles_math_connectors_unchanged=True,changed_pixels=int(delta.sum()),pixels_changed_outside_approved_text_boxes=outside,filter_or_filtered_visible_occurrences=0,clip_EMA_unchanged=True)
 (FINAL/f'Figure{n}_object_changes.json').write_text(json.dumps(proof,indent=2));checks.append(proof)
 linear=np.asarray(im,dtype=float)/255;linear=np.where(linear<=.04045,linear/12.92,((linear+.055)/1.055)**2.4)
 for name,m in {'deuteranopia':[[.367322,.860646,-.227968],[.280085,.672501,.047413],[-.011820,.042940,.968881]],'protanopia':[[.152286,1.052583,-.204868],[.114503,.786281,.099216],[-.003882,-.048116,1.051998]]}.items():
  a=np.clip(linear@np.array(m).T,0,1);a=np.where(a<=.0031308,12.92*a,1.055*a**(1/2.4)-.055);Image.fromarray(np.round(a*255).astype('uint8')).save(FINAL/(stem+'_'+name+'.png'),dpi=(300,300))
 print(json.dumps(proof))

shutil.copy2(SOURCE/'Figure2_loop_geometry.json',FINAL/'Figure2_loop_geometry.json')
shutil.copy2(SOURCE/'legend_measurement.json',FINAL/'legend_measurement.json')
assert baseline_hashes=={p.name:sha(p) for p in SOURCE.iterdir() if p.is_file()}
(FINAL/'checks.json').write_text(json.dumps(checks,indent=2));(FINAL/'requirements.json').write_text(json.dumps({'status':'PASS','changed_labels':4,'unchanged_layout_math_styles':True,'outside_change_box_pixel_differences':0,'r3_files_unchanged':baseline_hashes,'cold_review':'pending'},indent=2))
(FINAL/'HANDOFF.md').write_text('''# F02 R4: terminology only

R4 was made from exact copies of the final R3 PPTXs, not from the separate attached-PDF branch. Only four native text nodes changed:

- Figure1: “hard event: excluded edge” → “hard constraint: excluded edge”. “Fixed penalty” and “Lagrangian” are unchanged.
- Figure2: “Soft budgets” → “Soft constraints”; “Hard exclusions” → “Hard constraints”; “budget duals” → “safety prices”.
- No visible “filter” or “filtered” occurs in either R3 figure. The internal object ID `filter_chain` is unchanged, and its visible “clip + EMA” equation is unchanged.

All other PPTX ZIP entries are byte-identical to R3. Within slide1.xml, reverting these four strings reproduces canonical R3 XML exactly. Positions, dimensions, fonts, styling, formulas, groups, native editability and connectors are unchanged. Actual LibreOffice 300dpi renders differ only within the four existing text boxes: zero changed pixels elsewhere. Each replacement fits its original box. Both figures remain5.5×3.3in; exported labels>=8.107pt and unchanged formula scripts>=8.169pt. Zero raster objects.

Deliverables: Figure1/2_F02_R4.pptx, .svg, .pdf, _300dpi.png; grayscale and CVD previews; unchanged LaTeX and formula audits; text-only object-change reports, checks and source ZIP. Export uses the same LibreOffice → anonymous PDF/SVG/PNG pipeline as R3. Formulas are editable vector outlines, not OfficeMath; PowerPoint rendering is unverified.

R3 and the author’s separate attached-PDF/current-architecture branch remain unchanged. R4 awaits cold review; manuscript source stays with the owner.
''')
(FINAL/'DELIVERY.md').write_text('F02 R4: native text-only terminology edits on R3 copies. See HANDOFF.md for exact changes, unchanged-object proof and rendering limits.\n')
(FINAL/'PROCESS_NOTES.md').write_text('One typography-only round. Four labels replaced in native a:t nodes without reserializing the other package parts. Exact XML reversal and 300dpi pixel masks prove no changes elsewhere. No math rebuild or layout changes were needed.\n')
shutil.copy2(Path(__file__),FINAL/'revise.py')
with zipfile.ZipFile(FINAL/'F02_R4_sources.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in sorted(FINAL.iterdir()):
  if p.suffix in ['.pptx','.tex','.json','.py','.md']:z.write(p,p.name)
