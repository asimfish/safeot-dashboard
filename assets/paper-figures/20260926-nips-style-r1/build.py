from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent
FONT = 'Times New Roman'
DARK = RGBColor(20, 40, 55)
BLUE = RGBColor(24, 24, 230)
GREEN = RGBColor(18, 85, 41)
ORANGE = RGBColor(232, 139, 16)
GRAY = RGBColor(64, 64, 64)


def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if hasattr(sh, 'shapes'):
            yield from all_shapes(sh.shapes)


def find(prs, name):
    for slide in prs.slides:
        for sh in all_shapes(slide.shapes):
            if sh.name == name:
                return sh
    raise KeyError(name)


def set_text(prs, name, text, size=None, bold=None, color=None, align=None):
    sh = find(prs, name)
    if not hasattr(sh, 'text_frame'):
        raise TypeError(name)
    tf = sh.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align if align is not None else PP_ALIGN.CENTER
    run = p.add_run(); run.text = text
    f = run.font; f.name = FONT
    if size is not None: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color
    return sh


def replace_picture(prs, name, image_path):
    sh = find(prs, name)
    if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
        raise TypeError((name, sh.shape_type))
    image_part, rid = sh.part.get_or_add_image_part(str(image_path))
    blip = sh._element.blipFill.blip
    blip.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', rid)
    return sh


def remove_group(prs, name):
    sh = find(prs, name)
    if not hasattr(sh, 'shapes'):
        raise TypeError(name)
    el = sh._element
    el.getparent().remove(el)


def add_text(slide, x, y, w, h, text, size=8.4, color=DARK, bold=False,
             fill=None, line=None, align=PP_ALIGN.CENTER, margin=0.03,
             name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    if fill is not None:
        box.fill.solid(); box.fill.fore_color.rgb = fill
    else:
        box.fill.background()
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
    return box


def build_fig1(src, dst):
    prs = Presentation(str(src)); slide = prs.slides[0]
    # Keep the NIPS panel geometry, curves and spectrum. Replace only the narrative labels.
    set_text(prs, 'left_title', 'Penalty view', size=10.8, bold=True, color=BLUE)
    set_text(prs, 'left_st', 'costs → objective', size=6.3, bold=False, color=GRAY)
    set_text(prs, 'leakage', 'leakage', size=6.4, color=RGBColor(220, 20, 24))
    set_text(prs, 'left_weakness', 'weakness: fixed-price leakage', size=6.5, color=BLUE)
    set_text(prs, 'middle_title', 'SafeOT', size=12.5, bold=True, color=GREEN)
    set_text(prs, 'constraints_title', '1. Constraints', size=7.3, bold=True, color=GREEN)
    set_text(prs, 'compress_title', '2. One flow', size=7.3, bold=True, color=GREEN)
    set_text(prs, 'solve_title', '3. Actor', size=7.3, bold=True, color=GREEN)
    set_text(prs, 'sinkhorn_label', '(entropic)', size=7.0, color=GREEN)
    # Compact symbols keep the card rhythm without experiment-looking numeric tables.
    for i, (nm, val) in enumerate([
        ('C₁', 'C₁ · b₁'), ('C₂', 'C₂ · b₂'), ('C₃', 'C₃ · b₃'), ('H', 'H: F_H = 0')]):
        set_text(prs, f'costname{i}', nm, size=7.0, bold=True, color=[BLUE, RGBColor(27,169,192), GRAY, RGBColor(220,20,24)][i])
        set_text(prs, f'costvalues{i}', val, size=7.0, color=[BLUE, RGBColor(27,169,192), GRAY, RGBColor(220,20,24)][i])
    set_text(prs, 'compression_name', 'one flow', size=7.0, bold=True, color=GREEN)
    set_text(prs, 'uniform_min', 'F★ → λ★', size=7.0, bold=True, color=GREEN)
    set_text(prs, 'solvetext', 'λ★ + β\n→ actor', size=7.0, bold=True, color=GREEN)
    for nm, val in zip(['Pgrid_v0_0','Pgrid_v0_1','Pgrid_v0_2','Pgrid_v1_0','Pgrid_v1_1','Pgrid_v1_2','Pgrid_v2_0','Pgrid_v2_1','Pgrid_v2_2'],
                       ['λ₁','λ₂','λ₃','b₁','b₂','b₃','F★','β','λ']):
        set_text(prs, nm, val, size=5.7, bold=True, color=GREEN)
    set_text(prs, 'transport', 'safety prices', size=7.0, bold=True, color=GREEN)
    set_text(prs, 'hard_title', 'Constraint', size=10.8, bold=True, color=RGBColor(220,20,24))
    set_text(prs, 'hard_weak', 'weakness: late / oscillation', size=6.5, color=RGBColor(220,20,24))
    set_text(prs, 'soft_end', 'penalty', size=6.0, color=BLUE)
    set_text(prs, 'hard_end', 'constraints', size=6.0, color=RGBColor(220,20,24))
    set_text(prs, 'marktext0', 'fixed', size=7.0, color=GRAY)
    set_text(prs, 'marktext1', 'Lagrangian', size=7.0, color=GRAY)
    set_text(prs, 'marktext2', 'SafeOT', size=7.0, bold=True, color=GREEN)
    set_text(prs, 'marktext3', 'LP ε→0 / active', size=6.5, color=GRAY)
    set_text(prs, 'spectrum_name', 'one price rule', size=8.0, bold=True, color=RGBColor(255,255,255))
    try: remove_group(prs, 'middle_sub')
    except KeyError: pass
    # Generated pictograms replace the legacy tiny raster icons while preserving their card geometry.
    for nm, fn in zip(['costicon0','costicon1','costicon2','costicon3'], ['icon_graph.png','icon_price.png','icon_reroute.png','icon_actor.png']):
        replace_picture(prs, nm, ROOT / fn)
    prs.save(str(dst))


def build_fig2(src, dst):
    prs = Presentation(str(src)); slide = prs.slides[0]
    set_text(prs, 'title1', 'Rollout', size=9.6, bold=True, color=DARK)
    set_text(prs, 'collect', 'Collect transitions\n(sₜ, aₜ, rₜ, cₖ,ₜ, sₜ₊₁)', size=7.0, color=DARK)
    set_text(prs, 'title2', 'Flow inputs', size=9.6, bold=True, color=DARK)
    set_text(prs, 'graph_title', 'Rollout graph', size=7.2, color=DARK)
    set_text(prs, 'title3', 'Constraints', size=9.6, bold=True, color=GREEN)
    set_text(prs, 'one_matrix', 'soft: bₖ = Bₖ/L\nhard: F_H = 0', size=7.0, color=GREEN)
    set_text(prs, 'title4', 'Flow projection', size=9.2, bold=True, color=GREEN)
    set_text(prs, 'iterations', 'F̂ → F★  |  λ★', size=7.0, bold=True, color=GREEN)
    set_text(prs, 'iter_row', 'flow', size=7.0, color=GREEN)
    set_text(prs, 'iter_col', 'cost', size=7.0, color=GREEN)
    set_text(prs, 'iter_projected', 'target F★', size=7.0, color=GREEN)
    set_text(prs, 'repeat_label', 'target flow is not executed', size=7.0, color=GREEN)
    set_text(prs, 'title5', 'Priced PPO', size=8.6, bold=True, color=ORANGE)
    set_text(prs, 'weights', 'λ = λ̄★ + β', size=7.3, bold=True, color=ORANGE)
    set_text(prs, 'ppo_label', 'priced advantage', size=6.8, bold=True, color=ORANGE)
    # Remove stale solver/importance-weight equations and labels, retain colorful matrices/boxes.
    for nm in ['projection_eq', 'composeeq', 'weight_eq', 'ppo_eq1', 'ppo_eq2', 'emp_label', 'score_label', 'cap_label', 'tuple', 'cert', 'updated_policy']:
        try: remove_group(prs, nm)
        except KeyError: pass
    # Add concise current equations over the original NIPS boxes, preserving the visual rhythm.
    add_text(slide, 3.18, 0.69, 1.10, 0.18, 'F★ = projε(F̂; b, H)', size=7.0, color=GREEN, bold=True, fill=RGBColor(255,255,255), line=GREEN, name='current_projection_eq')
    add_text(slide, 4.56, 0.92, 0.80, 0.16, 'λ = λ̄★ + β', size=7.0, color=ORANGE, bold=True, fill=RGBColor(255,247,238), line=ORANGE, name='current_price_eq')
    add_text(slide, 4.53, 1.57, 0.84, 0.30, 'Ã = Aᵣ − λ · A𝑐', size=7.0, color=ORANGE, bold=True, fill=RGBColor(255,247,238), line=ORANGE, name='current_advantage_eq')
    add_text(slide, 4.52, 1.82, 0.84, 0.24, 'PPO / TRPO\nπₜ₊₁', size=7.0, color=ORANGE, bold=True, fill=RGBColor(255,247,238), line=ORANGE, name='current_update_eq')
    add_text(slide, 3.34, 2.52, 0.74, 0.12, 'safety prices λ★', size=7.0, color=GREEN, bold=True, fill=RGBColor(240,248,240), line=GREEN, name='current_safety_prices')
    add_text(slide, 1.37, 1.62, 0.60, 0.10, 'F̂ · A · Cₖ', size=7.0, color=BLUE, bold=True, name='current_input_key')
    # Use generated pictograms for the rollout and actor anchors; keep NIPS-native matrix artwork.
    for nm, fn in [('world','icon_rollout.png'), ('robot','icon_graph.png'), ('robot_next','icon_actor.png')]:
        try: replace_picture(prs, nm, ROOT / fn)
        except KeyError: pass
    prs.save(str(dst))


if __name__ == '__main__':
    build_fig1(ROOT / 'Figure1_NIPSStyle_R1.pptx', ROOT / 'Figure1_NIPSStyle_R1.pptx')
    build_fig2(ROOT / 'Figure2_NIPSStyle_R1.pptx', ROOT / 'Figure2_NIPSStyle_R1.pptx')
    print('built', ROOT)
